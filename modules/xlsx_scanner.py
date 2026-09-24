"""Batch Excel enrichment.

Detects the website column in any Excel (.xlsx) file, scrapes every listed
website for phone numbers and e-mails, and writes an enriched copy of the
original workbook as ``<input_filename>-scanned.xlsx`` with the new
``phone-scan`` and ``mail-scan`` columns appended.
"""

import re
import sys
import threading
from argparse import Namespace
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import requests

_PHONE_SCAN_COL = "phone-scan"
_MAIL_SCAN_COL = "mail-scan"

# Normalized (lowercase, non-alphanumeric stripped) column-name keywords that
# identify a "website" column. English + common Turkish variants are covered.
_COLUMN_KEYWORDS = {
    "website", "websitesi", "websi̇te", "websiteurl", "websitelink", "websitelink1",
    "websiteaddress", "businesswebsite", "companywebsite",
    "url", "urls", "link", "links", "domain", "domains",
    "homepage", "site", "sitesi", "sitelink", "siteurl", "siteaddress", "siteadresi",
    "web", "webadresi", "webaddress", "weblink", "home", "websiteadresi", "weburl",
}


def _normalize_header(header) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(header).lower())


def _looks_like_url(value) -> bool:
    if value is None:
        return False
    s = str(value).strip()
    if not s:
        return False
    lower = s.lower()
    if "://" in lower or lower.startswith(("http", "www.")):
        return True
    host = lower.split("/")[0]
    return "." in host and " " not in host and "@" not in host and not host.endswith(".")


def detect_website_column(headers: list, rows: list) -> str | None:
    """Pick the column that most likely holds website URLs.

    Keyword matches are strongly preferred; otherwise the column containing
    the most URL-like values is chosen. Returns ``None`` when nothing fits.
    """
    best, best_score = None, -1
    for idx, header in enumerate(headers):
        url_count = sum(1 for row in rows if _looks_like_url(row[idx]))
        if _normalize_header(header) in _COLUMN_KEYWORDS:
            url_count += 1000
        if url_count > best_score:
            best, best_score = header, url_count
    return best if best_score > 0 else None


def scan_xlsx(input_path: str, args: Namespace, verbose) -> Path:
    """Enrich an Excel file with scraped contact data and return the output path."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("openpyxl is required for --scan. Install it with: pip install openpyxl")

    # Lazy import avoids a circular import at module load time.
    from TheScrapper import normalize_url, scrape

    src = Path(input_path)
    if not src.exists():
        raise SystemExit(f"File not found: {src}")

    wb = load_workbook(str(src))
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    rows = [[cell.value for cell in row] for row in ws.iter_rows(min_row=2)]

    website_col = args.scan_column or None
    if website_col is not None:
        if website_col not in headers:
            raise SystemExit(
                f"Column '{website_col}' not found in {src.name}. "
                f"Available columns: {', '.join(str(h) for h in headers if h)}"
            )
    else:
        website_col = detect_website_column(headers, rows)
        if website_col is None:
            raise SystemExit(
                "Could not auto-detect a website column in "
                f"{src.name}. Available columns: "
                f"{', '.join(str(h) for h in headers if h)}. "
                "You can force it with --scan-column NAME."
            )
        verbose(f"Auto-detected website column: '{website_col}'")

    col_idx = headers.index(website_col)

    total = len(rows)
    urls = []
    row_by_url = {}
    for i, row in enumerate(rows, start=2):
        val = row[col_idx] if col_idx < len(row) else None
        if val and str(val).strip():
            raw = str(val).strip()
            url = normalize_url(raw)
            urls.append(url)
            row_by_url[url] = i

    if not urls:
        raise SystemExit(f"No website URLs found in column '{website_col}'.")

    total = len(urls)

    verbose(f"Scraping {len(urls)} websites with {args.threads} threads...")

    scan_args = Namespace(
        email=True,
        number=True,
        socials=False,
        social_extract=False,
        crawl=args.crawl,
        crawl_external=args.crawl_external,
        threads=args.threads,
    )

    # Append the new columns.
    phone_col = ws.max_column + 1
    mail_col = ws.max_column + 2
    ws.cell(row=1, column=phone_col, value=_PHONE_SCAN_COL)
    ws.cell(row=1, column=mail_col, value=_MAIL_SCAN_COL)

    lock = threading.Lock()
    completed = 0
    failed = 0

    def _scrape_one(url):
        try:
            return url, scrape(url, scan_args)
        except requests.exceptions.RequestException:
            return url, None

    with ThreadPoolExecutor(max_workers=args.threads) as pool:
        futures = {pool.submit(_scrape_one, url): url for url in urls}
        for future in as_completed(futures):
            url, result = future.result()
            with lock:
                completed += 1
                row_idx = row_by_url[url]
                if result is not None:
                    ws.cell(row=row_idx, column=phone_col, value="; ".join(result.get("Numbers", [])))
                    ws.cell(row=row_idx, column=mail_col, value="; ".join(result.get("E-Mails", [])))
                else:
                    failed += 1
                sys.stderr.write(
                    f"\r[{completed}/{total}] Done: {completed - failed} | Failed: {failed}  "
                )
                sys.stderr.flush()

    sys.stderr.write("\n")

    Path("output").mkdir(exist_ok=True)
    out_path = Path("output") / f"{src.stem}-scanned.xlsx"
    wb.save(out_path)
    wb.close()
    return out_path